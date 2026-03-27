import json
import math
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Pt, Inches, Cm
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR
import os
import re

class PPTGenerator:
    def __init__(self, template_path: str, json_path: str):
        self.prs = Presentation(template_path)

        with open(json_path, "r", encoding="utf-8") as f:
            self.data = json.load(f)

        # Template có 41 slides (index 0–40):
        #   [0]      → add_cover_slide()        (edit in-place)
        #   [1]      → add_table_of_content()   (edit in-place)
        #   [2..39]  → content slides           (lấy tuần tự qua _next_slide())
        #   [40]     → add_end_slide()          (edit in-place)

        # Con trỏ slide content — bắt đầu từ index 2
        self._slide_cursor = 2

        # Màu chủ đạo
        self.primary_red = RGBColor(0xB5, 0x1F, 0x09)

    # ================================================================
    # UTIL
    # ================================================================

    def _next_slide(self):
        """Trả về slide tiếp theo trong vùng content (index 2 → 39), tự tăng cursor."""
        idx = self._slide_cursor
        if idx > 39:
            raise IndexError(
                f"Đã dùng hết slide template content (cursor={idx}). "
                "Tăng số slide trong template hoặc giảm dữ liệu."
            )
        self._slide_cursor += 1
        return self.prs.slides[idx]

    def _get_all_shapes(self, shapes):
        all_shapes = []
        for shape in shapes:
            all_shapes.append(shape)
            if shape.shape_type == 6:  # GROUP
                for sub in self._get_all_shapes(shape.shapes):
                    all_shapes.append(sub)
        return all_shapes

    def _find_shape_by_name(self, shapes, name: str):
        for shape in shapes:
            if shape.name == name:
                return shape
        return None

    def hex_to_rgb(self, hex_color: str):
        hex_color = hex_color.lstrip("#")
        return tuple(int(hex_color[i:i+2], 16) for i in (0, 2, 4))
    
    def hex_to_rgb_color(self, hex_color: str):
        hex_color = hex_color.lstrip("#")
        r, g, b = tuple(int(hex_color[i:i+2], 16) for i in (0, 2, 4))
        return RGBColor(r, g, b)

    def _bring_to_front(self, shape):
        sp = shape._element
        parent = sp.getparent()
        parent.remove(sp)
        parent.append(sp)

    def _set_text_exact_style(
        self, shape, text: str,
        font_name: str = "Arial", font_size: int = 40,
        bold: bool = True, italic: bool = False,
        color: str | None = None, align: str = "center"
    ):
        if not shape or not shape.has_text_frame:
            return
        tf = shape.text_frame
        tf.clear()
        p = tf.paragraphs[0]
        align_map = {
            "left":    PP_ALIGN.LEFT,
            "center":  PP_ALIGN.CENTER,
            "right":   PP_ALIGN.RIGHT,
            "justify": PP_ALIGN.JUSTIFY,
        }
        p.alignment = align_map.get(align.lower(), PP_ALIGN.CENTER)
        run = p.add_run()
        run.text = text
        font = run.font
        font.name = font_name
        font.size = Pt(font_size)
        font.bold = bold
        font.italic = italic
        if color:
            rgb = self.hex_to_rgb(color)
            font.color.rgb = RGBColor(*rgb)

    def add_text(
        self, slide, text: str,
        left: float, top: float, width: float, height: float,
        font_size: int = 24,
        bold: bool = False,
        color=None,
        align=PP_ALIGN.LEFT,
        italic: bool = False,
        font_name: str = "Arial"
    ):
        """
        Thêm textbox lên slide.
        color: str hex ("B51F09") hoặc RGBColor.
        align: PP_ALIGN enum hoặc str ("left"/"center"/"right").
        """
        box = slide.shapes.add_textbox(
            Inches(left), Inches(top), Inches(width), Inches(height)
        )
        tf = box.text_frame
        tf.word_wrap = True
        tf.clear()

        align_map = {
            "left":    PP_ALIGN.LEFT,
            "center":  PP_ALIGN.CENTER,
            "right":   PP_ALIGN.RIGHT,
            "justify": PP_ALIGN.JUSTIFY,
        }
        if isinstance(align, str):
            align = align_map.get(align.lower(), PP_ALIGN.LEFT)

        for idx, line in enumerate(str(text).split("\n")):
            para = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
            para.alignment = align
            run = para.add_run()
            run.text        = line
            run.font.size   = Pt(font_size)
            run.font.bold   = bold
            run.font.italic = italic
            run.font.name   = font_name
            if color is not None:
                if isinstance(color, RGBColor):
                    run.font.color.rgb = color
                elif isinstance(color, str):
                    run.font.color.rgb = RGBColor(*self.hex_to_rgb(color))

        self._bring_to_front(box)
        return box

    def get_text_height(self, text: str, font_size: int, width_inches: float) -> float:
        """Ước tính chiều cao (inches) cho một đoạn text, hỗ trợ CJK full-width."""
        SLIDE_HEIGHT   = 7.5
        width_pt       = width_inches * 72
        char_width_pt  = font_size * 0.6
        units_per_line = width_pt / char_width_pt

        def char_units(ch):
            cp = ord(ch)
            if (0x4E00 <= cp <= 0x9FFF or
                    0x3000 <= cp <= 0x303F or
                    0xFF00 <= cp <= 0xFFEF or
                    0x3040 <= cp <= 0x30FF):
                return 1.0
            return 0.5

        total_lines = 0
        for raw_line in str(text).split("\n"):
            if raw_line.strip() == "":
                total_lines += 1
                continue
            line_units = sum(char_units(c) for c in raw_line)
            total_lines += max(1, math.ceil(line_units / units_per_line))

        line_height_pt  = font_size * 1.2
        total_height_pt = total_lines * line_height_pt + 8
        return (total_height_pt / (SLIDE_HEIGHT * 72)) * SLIDE_HEIGHT

    def _apply_title_styling(self, shape, text: str, max_font=135, min_font=40):
        """
        Căn chỉnh font-size tự động để title dài vừa vặn trên 1 dòng
        mà không bị vượt quá chiều rộng và chiều cao của khung Textbox.
        """
        if not text:
            return

        # Ép chữ nằm trên 1 dòng
        tf = shape.text_frame
        tf.word_wrap = False
        
        # 1 point (pt) = 12700 EMUs trong thư viện python-pptx
        # Tính toán kích thước thực tế của text box (theo đơn vị Point)
        shape_width_pt = shape.width / 12700.0
        shape_height_pt = shape.height / 12700.0
        
        # Lấy thông số lề (margin) của Textbox. Nếu PPT không set thì lấy giá trị mặc định
        margin_left = tf.margin_left / 12700.0 if tf.margin_left is not None else 7.2
        margin_right = tf.margin_right / 12700.0 if tf.margin_right is not None else 7.2
        margin_top = tf.margin_top / 12700.0 if tf.margin_top is not None else 3.6
        margin_bottom = tf.margin_bottom / 12700.0 if tf.margin_bottom is not None else 3.6
        
        # Không gian thực tế an toàn được phép chứa chữ (đã trừ lề)
        available_width_pt = shape_width_pt - margin_left - margin_right
        available_height_pt = shape_height_pt - margin_top - margin_bottom
        
        # Tính "đơn vị" bề rộng (1 chữ Hán/Full-width = 1 unit, 1 chữ Latinh/Số = 0.55 unit)
        units = 0.0
        for ch in text:
            cp = ord(ch)
            if (0x4E00 <= cp <= 0x9FFF or 0x3000 <= cp <= 0x303F or 
                0xFF00 <= cp <= 0xFFEF or 0x3040 <= cp <= 0x30FF):
                units += 1.0
            else:
                units += 0.55

        if units == 0:
            optimal_font = max_font
        else:
            # 1. Ràng buộc Chiều rộng: Size chữ tối đa để dàn đủ số ký tự trên bề ngang
            max_font_by_width = (available_width_pt / units) * 0.95 # Nhân 0.95 để trừ hao an toàn 5%
            
            # 2. Ràng buộc Chiều cao: Size chữ tối đa không được lớn hơn chiều cao Textbox
            max_font_by_height = available_height_pt * 0.90 # Trừ hao 10% cho khoảng cách dòng (line spacing)
            
            # Lấy size chữ thoả mãn CẢ 2 ĐIỀU KIỆN (Không tràn ngang, không tràn dọc)
            calculated_font = min(max_font_by_width, max_font_by_height)
            
            # Chốt font size cuối cùng nằm trong khoảng giới hạn cho phép
            optimal_font = int(min(max_font, max(min_font, calculated_font)))

        self._set_text_exact_style(
            shape, text, font_size=optimal_font, color="B51F09", align="center"
        )

    def _replace_subtitle_with_rounded_rect(self, slide, old_shape, text: str, font_size: int = 50):
        """Thay thế textbox cũ bằng khối bo tròn màu đỏ có chữ trắng căn giữa."""
        
        center_x = old_shape.left + old_shape.width / 2.0
        center_y = old_shape.top + old_shape.height / 2.0
        
        # Ước tính chiều rộng cần thiết: mỗi ký tự tốn khoảng (0.4 * font_size) points
        # Cộng thêm padding ~ 1.5 inches để nhìn cân đối
        est_width_in = len(text) * (font_size * 0.4 / 72) + 1.5
        new_width = int(Inches(est_width_in))
        new_height = int(Inches(1.2))  # Chiều cao cố định (1.2") đủ rộng cho text 50pt
        
        new_left = int(center_x - new_width / 2.0)
        new_top = int(center_y - new_height / 2.0)
        
        # Xóa shape cũ (ẩn đi bằng cách gỡ khỏi XML element tree)
        old_element = old_shape._element
        old_element.getparent().remove(old_element)
        
        # Tạo shape bo tròn mới
        new_shape = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE, new_left, new_top, new_width, new_height
        )
        
        # Pill shape (bo tròn hoàn toàn hai đầu)
        new_shape.adjustments[0] = 0.5 
        
        # Định dạng màu đỏ cho fill và tắt viền
        new_shape.fill.solid()
        new_shape.fill.fore_color.rgb = RGBColor(0xB5, 0x1F, 0x09)
        new_shape.line.fill.background()
        
        # Căn giữa chữ theo chiều dọc và ngang
        tf = new_shape.text_frame
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        
        self._set_text_exact_style(
            new_shape, text, font_size=font_size, color="FFFFFF", align="center"
        )
        self._bring_to_front(new_shape)

    def _add_hl_text(self, paragraph, text, default_font, default_size, default_color, hl_color="29741D"):
        """
        Hàm hỗ trợ bóc tách thẻ <hl>...</hl> và add vào paragraph với định dạng khác nhau.
        """
        # Tách chuỗi dựa trên thẻ <hl> và </hl>
        parts = re.split(r'(<hl>.*?</hl>)', text)
        
        for part in parts:
            run = paragraph.add_run()
            if part.startswith('<hl>') and part.endswith('</hl>'):
                # Phần nằm trong thẻ highlight: In đậm và màu xanh lá
                content = part[4:-5]
                run.text = content
                run.font.bold = True
                run.font.color.rgb = self.hex_to_rgb_color(hl_color)
            else:
                # Phần text bình thường: Không đậm và màu mặc định
                run.text = part
                run.font.bold = False
                run.font.color.rgb = self.hex_to_rgb_color(default_color)
            
            run.font.name = default_font
            run.font.size = Pt(default_size)

    # ================================================================
    # FIXED SLIDES — edit in-place
    # ================================================================

    def add_cover_slide(self):
        """Slide [0] — edit in-place."""
        lesson = self.data.get("lesson_info", {})
        slide  = self.prs.slides[0]
        shapes = self._get_all_shapes(slide.shapes)

        title_shape    = self._find_shape_by_name(shapes, "TextBox 18")
        subtitle_shape = self._find_shape_by_name(shapes, "TextBox 30")

        if title_shape:
            # Chiều cao box template của bạn là ~2.78 inches (7.06cm)
            self._apply_title_styling(
                title_shape, lesson.get("title_cn", ""),
                max_font=135, min_font=75
            )
        if subtitle_shape:
            self._replace_subtitle_with_rounded_rect(
                slide, subtitle_shape,
                f"Bài Khóa - Bài {lesson.get('number', '')}",
                font_size=50
            )

    def add_table_of_content(self):
        """Slide [1] — edit in-place."""
        slide  = self.prs.slides[1]
        shapes = self._get_all_shapes(slide.shapes)

        left_shape  = self._find_shape_by_name(shapes, "TextBox 16")
        title_shape = self._find_shape_by_name(shapes, "TextBox 17")
        right_shape = self._find_shape_by_name(shapes, "TextBox 18")

        if left_shape:
            self._set_text_exact_style(left_shape,  "01. TỪ VỰNG", "Anton",  font_size=56, color="FCF1D4")
        if title_shape:
            self._set_text_exact_style(title_shape, "目录",                 font_size=90, color="B02012")
        if right_shape:
            self._set_text_exact_style(right_shape, "02. BÀI KHÓA", "Anton", font_size=56, color="FCF1D4")

    def add_end_slide(self):
        """Slide [40] — edit in-place."""
        lesson = self.data.get("lesson_info", {})
        slide  = self.prs.slides[40]
        shapes = self._get_all_shapes(slide.shapes)

        title_shape    = self._find_shape_by_name(shapes, "TextBox 10")
        subtitle_shape = self._find_shape_by_name(shapes, "TextBox 21")

        if title_shape:
            # Tương tự như slide cover, sử dụng min_font là 75
            self._apply_title_styling(
                title_shape, lesson.get("title_cn", ""),
                max_font=126, min_font=75
            )
        if subtitle_shape:
            self._replace_subtitle_with_rounded_rect(
                slide, subtitle_shape,
                f"Bài Khóa - Bài {lesson.get('number', '')}",
                font_size=50
            )

    # ================================================================
    # CONTENT SLIDES — Lập trình các hàm chèn nội dung vào đây
    # ================================================================

    def add_dialogue_slide(self, sec: dict, force_simple: bool = False):
        """
        Sinh slide Hội thoại. 
        Tự động điều chỉnh Line Spacing linh hoạt để Case 6 câu không bị rời rạc.
        """
        if force_simple:
            dialogue_data = sec.get("simple_dialogue")
            has_pinyin = False
        else:
            dialogue_data = sec.get("full_dialogue")
            has_pinyin = True

        if not dialogue_data: return  

        hz_list = dialogue_data.get("hz", [])
        py_list = dialogue_data.get("py", []) if has_pinyin else []
        vi_list = dialogue_data.get("vi", [])
        
        num_sentences = len(hz_list)
        if num_sentences == 0: return

        MAX_PER_SLIDE = 6
        chunks = [ (hz_list[i:i+6], py_list[i:i+6], vi_list[i:i+6]) for i in range(0, num_sentences, 6)]

        flower_icon_path = r"D:\Edmicro\Tools\create_hsk\hsk_ppt\images\flower_point.png"
        section_title = sec.get("section_title", "Hội thoại")

        for chunk_idx, (chunk_hz, chunk_py, chunk_vi) in enumerate(chunks):
            slide = self._next_slide()
            num_items = len(chunk_hz)

            # A. TITLE
            title_width, title_height = Cm(20.14), Cm(2.67)
            title_left, title_top = (self.prs.slide_width - title_width) / 2, Cm(1.26)
            title_box = slide.shapes.add_textbox(title_left, title_top, title_width, title_height)
            title_box.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
            self._set_text_exact_style(title_box, section_title, font_name="Fraunces", font_size=57, color="FCF1D4", align="center")

            # B. CONFIG DYNAMIC SPACING
            slide_width_cm = self.prs.slide_width / 360000.0 
            
            if num_items >= 5: 
                # --- CASE A: 5-6 CÂU THOẠI (Dùng spacing nhỏ hơn để gọn gàng) ---
                start_left, start_top = 7.95, 5.56
                box_width = slide_width_cm - start_left - 1.0 
                box_height = 3.19  
                hz_size, hz_font = 29.3, "字由点字典楷"  
                py_size, py_font = 18.2, "Muli"         
                vi_size, vi_font, vi_italic = 18.2, "Muli", False
                y_step = 3.7 
                
                # Tự động tính Spacing gọn hơn (ví dụ: 28.5pt cho font 29.3pt)
                line_spacing_val = Pt(28.5) 
            else: 
                # --- CASE B: 1-4 CÂU THOẠI (Dùng spacing 33pt như bạn thấy phù hợp) ---
                start_left, start_top = 11.95, 7.24
                box_width = slide_width_cm - start_left - 1.5
                box_height = 4.27  
                hz_size, hz_font = 39.5, "字由点字典楷" 
                py_size, py_font = 23.5, "Muli"         
                vi_size, vi_font, vi_italic = 24.5, "Muli", True
                y_step = 4.72 
                
                line_spacing_val = Pt(33) 

            # C. RENDER
            # Tính toán line_spacing theo cm để đặt icon chính xác
            line_spacing_cm = (line_spacing_val.pt / 72.0) * 2.54

            for i in range(num_items):
                current_top = start_top + (i * y_step)
                
                # Tính icon_top bám theo dòng Hán tự
                offset = line_spacing_cm if has_pinyin else (line_spacing_cm * 0.5)
                icon_top_val = current_top + (box_height / 2.0) - offset - 0.4
                
                if os.path.exists(flower_icon_path):
                    slide.shapes.add_picture(flower_icon_path, Cm(start_left - 1.2), Cm(icon_top_val), Cm(0.8), Cm(0.8))

                textbox = slide.shapes.add_textbox(Cm(start_left), Cm(current_top), Cm(box_width), Cm(box_height))
                tf = textbox.text_frame
                tf.word_wrap = True
                tf.vertical_anchor = MSO_ANCHOR.MIDDLE
                tf.margin_bottom = tf.margin_top = tf.margin_left = tf.margin_right = 0
                
                # Render các paragraph
                def add_para(text, font, size, color, is_bold=False, is_italic=False):
                    p = tf.add_paragraph() if len(tf.paragraphs[0].runs) > 0 else tf.paragraphs[0]
                    p.space_before = p.space_after = Pt(0)
                    p.line_spacing = line_spacing_val
                    run = p.add_run()
                    run.text = text
                    run.font.name, run.font.size, run.font.bold, run.font.italic = font, Pt(size), is_bold, is_italic
                    run.font.color.rgb = self.hex_to_rgb_color(color)

                add_para(chunk_hz[i], hz_font, hz_size, "000000", is_bold=True)
                if has_pinyin: add_para(chunk_py[i], py_font, py_size, "545454")
                add_para(chunk_vi[i], vi_font, vi_size, "A40400", is_italic=vi_italic)



    def add_vocab_slides(self, sec: dict):
        """
        Sinh các slide Từ vựng. 3 từ vựng / 1 slide.
        Dữ liệu lấy từ slide content của template.
        """
        vocab_list = sec.get("vocabulary", [])
        if not vocab_list:
            return

        # 1. Chia cụm 3 từ / slide
        MAX_PER_SLIDE = 3
        chunks = [vocab_list[i : i + MAX_PER_SLIDE] for i in range(0, len(vocab_list), MAX_PER_SLIDE)]

        for chunk in chunks:
            # Lấy slide content tiếp theo từ template
            slide = self._next_slide()

            # ==========================================
            # A. TEXTBOX TITLE "Từ vựng"
            # ==========================================
            title_width = Cm(20.14)
            title_height = Cm(2.67)
            title_left = Cm(15.28)
            title_top = Cm(1.26)
            
            title_box = slide.shapes.add_textbox(title_left, title_top, title_width, title_height)
            title_box.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
            self._set_text_exact_style(
                title_box, "Từ vựng", font_name="Fraunces", font_size=57, 
                color="FCF1D4", bold=False, align="center"
            )

            # ==========================================
            # B. THÔNG SỐ VỊ TRÍ & BƯỚC NHẢY
            # ==========================================
            # Thông số Word Box (Cột trái)
            word_left = 2.86
            word_start_top = 6.71
            word_width = 8.69
            word_height = 5.54
            word_y_step = 6.73 # 13.44 - 6.71

            # Thông số Detail Box (Cột phải)
            detail_left = 15.28
            detail_start_top = 7.30
            # Chiều rộng động: slide_width - detail_left - lề phải (1cm)
            slide_width_cm = self.prs.slide_width / 360000.0
            detail_width = slide_width_cm - detail_left - 1.0
            detail_height = 4.59
            detail_y_step = 6.87 # 14.17 - 7.30

            # Màu sắc
            color_gray = "545454"
            color_red = "A40400"
            color_black = "000000"
            color_green = "29741D" # Màu xanh lá highlight

            # ==========================================
            # C. RENDER TỪNG TỪ VỰNG TRONG CHUNK
            # ==========================================
            for i, item in enumerate(chunk):
                # --- 1. Vẽ Word Box (Hán tự to bên trái) ---
                curr_word_top = word_start_top + (i * word_y_step)
                hz_text = item.get("hz", "")
                
                # Quyết định font size dựa trên độ dài từ
                hz_font_size = 116.5 if len(hz_text) <= 4 else 86
                
                word_box = slide.shapes.add_textbox(
                    Cm(word_left), Cm(curr_word_top), Cm(word_width), Cm(word_height)
                )
                word_box.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
                self._set_text_exact_style(
                    word_box, hz_text, font_name="字由点字云霆楷体", 
                    font_size=hz_font_size, color=color_black, bold=False, align="center"
                )

                # --- 2. Vẽ Detail Box (Pinyin, Nghĩa, Ví dụ bên phải) ---
                curr_detail_top = detail_start_top + (i * detail_y_step)
                detail_box = slide.shapes.add_textbox(
                    Cm(detail_left), Cm(curr_detail_top), Cm(detail_width), Cm(detail_height)
                )
                tf = detail_box.text_frame
                tf.word_wrap = True
                # Tắt margin để text sát lề chuẩn
                tf.margin_top = tf.margin_bottom = tf.margin_left = tf.margin_right = 0

                # Paragraph 1: /pinyin/ (loại từ): nghĩa
                p1 = tf.paragraphs[0]
                p1.space_after = Pt(0)
                
                # Run: Pinyin (Xám)
                run_py = p1.add_run()
                run_py.text = f"/{item.get('pinyin', '')}/ "
                run_py.font.name = "Muli"
                run_py.font.size = Pt(32)
                run_py.font.color.rgb = self.hex_to_rgb_color(color_gray)
                
                # Run: Loại từ & Nghĩa (Đỏ)
                run_type_vi = p1.add_run()
                type_str = f"({item.get('type', '')}): " if item.get('type') else ""
                run_type_vi.text = f"{type_str}{item.get('vi', '')}"
                run_type_vi.font.name = "Muli"
                run_type_vi.font.size = Pt(32)
                run_type_vi.font.color.rgb = self.hex_to_rgb_color(color_red)

                # Lấy dữ liệu ví dụ
                example = item.get("example", {})
                ex_hz = example.get("hz", "")
                ex_py = example.get("py", "")
                ex_vi = example.get("vi", "")

                # Paragraph 2: VD: Hán tự /Pinyin/
                if ex_hz or ex_py:
                    p2 = tf.add_paragraph()
                    p2.space_after = Pt(0)
                    
                    # Run: Chữ "VD: " (Đen, Size 32)
                    run_vd = p2.add_run()
                    run_vd.text = "VD: "
                    run_vd.font.name = "Muli"
                    run_vd.font.size = Pt(32)
                    run_vd.font.color.rgb = self.hex_to_rgb_color(color_black)
                    
                    # Chèn Hán tự ví dụ (Highlight thẻ <hl>) - Size 32
                    self._add_hl_text(p2, ex_hz, "Muli", 32, color_black, color_green)
                    
                    # Chèn Pinyin ví dụ bọc trong dấu / / - Phải gán size cho cả dấu gạch
                    if ex_py:
                        # Dấu "/" mở đầu
                        run_slash_start = p2.add_run()
                        run_slash_start.text = " /"
                        run_slash_start.font.name = "Muli"
                        run_slash_start.font.size = Pt(32)
                        run_slash_start.font.color.rgb = self.hex_to_rgb_color(color_black)

                        # Pinyin chính (có highlight) - Size 32
                        self._add_hl_text(p2, ex_py, "Muli", 32, color_black, color_green)

                        # Dấu "/:" kết thúc
                        run_slash_end = p2.add_run()
                        run_slash_end.text = "/:"
                        run_slash_end.font.name = "Muli"
                        run_slash_end.font.size = Pt(32)
                        run_slash_end.font.color.rgb = self.hex_to_rgb_color(color_black)

                # Paragraph 3: Dịch nghĩa ví dụ (Đã có size 32 trong hàm _add_hl_text)
                if ex_vi:
                    p3 = tf.add_paragraph()
                    p3.space_after = Pt(0)
                    self._add_hl_text(p3, ex_vi, "Muli", 32, color_black, color_green)

    def add_extra_slide(self, extra: dict):
        pass

    def add_exercise_slide(self, sec: dict):
        pass

    # ================================================================
    # BUILD & SAVE
    # ================================================================

    def build(self):
        """
        Build toàn bộ presentation theo thứ tự:
        [0] cover → [1] TOC → [2..39] dialogue/vocab/extra/exercise × n → [40] end
        """
        self.add_cover_slide()
        self.add_table_of_content()

        for sec in self.data["sections"]:
            # --- Phần 1: Hội thoại (Có Pinyin) ---
            self.add_dialogue_slide(sec, force_simple=False)
            
            # --- Phần 2: Từ vựng (3 từ/slide) ---
            self.add_vocab_slides(sec)
            
            # --- Phần 3: Kiến thức liên quan ---
            self.add_extra_slide(sec) 
            
            # --- Phần 4: Hội thoại (KHÔNG Pinyin) ---
            self.add_dialogue_slide(sec, force_simple=True)
            
            # --- Phần 5: Luyện tập ---
            self.add_exercise_slide(sec)

        # [TODO: Gọi các hàm vẽ nội dung (dialogue, vocab...) ở đây]

        self.add_end_slide()

        # Xóa các slide content thừa (từ _slide_cursor đến 39)
        # Lưu ý: Cần xóa ngược từ dưới lên để không làm sai lệch index của các slide còn lại
        for i in range(39, self._slide_cursor - 1, -1):
            rId = self.prs.slides._sldIdLst[i].rId
            self.prs.part.drop_rel(rId)
            del self.prs.slides._sldIdLst[i]

        used = self._slide_cursor - 2
        print(f"✅ Đã dùng {used} / 38 slide content "
              f"(index 2–{self._slide_cursor - 1}). "
              f"Đã xóa {38 - used} slide thừa.")

    def save(self, output_path: str):
        self.prs.save(output_path)
        print(f"✅ Saved: {output_path}")

# ================================================================
# MAIN
# ================================================================
if __name__ == "__main__":
    template_path = r"D:\Edmicro\Tools\create_hsk\hsk_ppt\template\HSK1 Bài Khóa template.pptx"
    json_path     = r"D:\Edmicro\Tools\create_hsk\hsk_ppt\test_output_json.json"
    output_path   = r"D:\Edmicro\Tools\create_hsk\hsk_ppt\test_output_baikhoa_new.pptx"

    gen = PPTGenerator(template_path, json_path)
    gen.build()
    gen.save(output_path)
