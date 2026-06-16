from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from services.callAPI import VertexClient
import mimetypes
import os
import json


FONT_NAME = "Times New Roman"
FONT_SIZE_TITLE = 18
FONT_SIZE_SECTION = 16
FONT_SIZE_BODY = 12


def _set_table_no_borders(table):
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)

    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "nil")


def _set_run_font(run, size=FONT_SIZE_BODY, bold=False, italic=False, color=None):
    run.font.name = FONT_NAME
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_NAME)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color is not None:
        run.font.color.rgb = color


def _set_paragraph_font(paragraph, size=FONT_SIZE_BODY, bold=False, italic=False):
    for run in paragraph.runs:
        _set_run_font(run, size=size, bold=bold, italic=italic)


def _set_document_default_font(doc):
    styles = doc.styles
    for style_name in ["Normal", "Title"]:
        style = styles[style_name]
        style.font.name = FONT_NAME
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_NAME)
        style.font.size = Pt(FONT_SIZE_BODY)


def _append_example_block(doc, example):
    if example.get("example_cn"):
        paragraph = doc.add_paragraph()
        run = paragraph.add_run(f"- {example['example_cn']}")
        _set_run_font(run)

    if example.get("example_pinyin"):
        paragraph = doc.add_paragraph()
        run = paragraph.add_run(example["example_pinyin"])
        _set_run_font(run)

    if example.get("example_vi"):
        paragraph = doc.add_paragraph()
        run = paragraph.add_run(example["example_vi"])
        _set_run_font(run, italic=True)


def response2docx_grammar(file_path, prompt, file_name, project_id, creds, model_name, vocab_records):
    client = VertexClient(project_id, creds, model_name)

    with open(file_path, "rb") as f:
        pdf_data = f.read()
    mime_type = mimetypes.guess_type(file_path)[0] or "application/pdf"

    schema_path = "resources/schema/summary_grammar_schema.json"
    if not os.path.exists(schema_path):
        raise FileNotFoundError(f"Không tìm thấy file schema tại {schema_path}")

    with open(schema_path, 'r', encoding='utf-8') as f:
        response_schema = json.load(f)

    ai_response_text = client.send_data_to_AI(
        prompt,
        data=pdf_data,
        mime_type=mime_type,
        response_mime_type="application/json",
        response_schema=response_schema,
        service_tier="flex"
    )

    try:
        data = json.loads(ai_response_text)
    except json.JSONDecodeError as e:
        raise ValueError(f"AI không trả về JSON hợp lệ: {e}") from e

    doc = Document()
    _set_document_default_font(doc)

    title = f"{data.get('lesson_number', '').strip()}: {data.get('lesson_title_vi', '').strip()}".strip(': ')
    if title:
        p_title = doc.add_paragraph()
        p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run_title = p_title.add_run(title)
        _set_run_font(run_title, size=FONT_SIZE_TITLE, bold=True)

    doc.add_paragraph()

    vocab_title = doc.add_paragraph()
    vocab_run = vocab_title.add_run("1. Từ vựng")
    _set_run_font(vocab_run, size=FONT_SIZE_SECTION, bold=True)

    table = doc.add_table(rows=1, cols=5)
    _set_table_no_borders(table)
    header_cells = table.rows[0].cells
    headers = ["STT", "Từ vựng", "Phiên âm", "Loại từ", "Nghĩa"]
    for idx, header in enumerate(headers):
        header_cells[idx].text = header
        for paragraph in header_cells[idx].paragraphs:
            for run in paragraph.runs:
                _set_run_font(run, bold=True)

    for row_idx, record in enumerate(vocab_records, 1):
        cells = table.add_row().cells
        cells[0].text = str(row_idx)
        cells[1].text = record.get("word", "")
        cells[2].text = record.get("pinyin", "")
        cells[3].text = record.get("part_of_speech", "")
        cells[4].text = record.get("meaning", "")
        for cell in cells:
            for paragraph in cell.paragraphs:
                _set_paragraph_font(paragraph)

    doc.add_paragraph()

    grammar_title = doc.add_paragraph()
    grammar_run = grammar_title.add_run("2. Ngữ pháp")
    _set_run_font(grammar_run, size=FONT_SIZE_SECTION, bold=True)

    for section in data.get("grammar_sections", []):
        p_section = doc.add_paragraph()
        run_section = p_section.add_run(f"Phần {section.get('section_number', '')}: {section.get('section_title', '')}")
        _set_run_font(run_section, italic=True)

        for item in section.get("items", []):
            p_item = doc.add_paragraph()
            run_item = p_item.add_run(item.get("highlight_text", ""))
            _set_run_font(run_item, color=RGBColor(0xC0, 0x00, 0x00))

            if item.get("detail_text"):
                p_detail = doc.add_paragraph()
                run_detail = p_detail.add_run(item.get("detail_text", ""))
                _set_run_font(run_detail)

            for example in item.get("examples", []):
                _append_example_block(doc, example)

        p_practice = doc.add_paragraph()
        practice_run = p_practice.add_run("Luyện tập")
        _set_run_font(practice_run, italic=True)

        practice = section.get("practice", {})
        practice_type = practice.get("practice_type", "")
        practice_label = "Trắc nghiệm" if practice_type == "multiple_choice" else "Điền từ"
        p_practice_type = doc.add_paragraph()
        practice_type_run = p_practice_type.add_run(practice_label)
        _set_run_font(practice_type_run)

        for idx, question in enumerate(practice.get("questions", []), 1):
            p_question = doc.add_paragraph()
            question_run = p_question.add_run(f"{idx}. {question.get('question_text', '')}")
            _set_run_font(question_run)

            options = question.get("options", [])
            for option in options:
                opt_p = doc.add_paragraph()
                opt_run = opt_p.add_run(str(option))
                _set_run_font(opt_run)

            if question.get("answer"):
                p_answer = doc.add_paragraph()
                run_answer = p_answer.add_run(f"Đáp án: {question.get('answer', '')}")
                _set_run_font(run_answer, italic=True)

        doc.add_paragraph()

    output_dir = "output/summary"
    os.makedirs(output_dir, exist_ok=True)
    output_path = os.path.join(output_dir, f"{file_name}_grammar_summary.docx")
    doc.save(output_path)
    return output_path