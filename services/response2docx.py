from docx import Document
from services.callAPI import VertexClient
from docx.shared import Inches
from io import BytesIO
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt
import mimetypes
import os
import json


def _append_note_examples(doc, examples):
    examples = examples or []
    if not examples:
        return

    p_examples_title = doc.add_paragraph()
    p_examples_title.add_run("Ví dụ:")

    for example in examples:
        example_cn = example.get('example_cn', '').strip()
        example_pinyin = example.get('example_pinyin', '').strip()
        example_vi = example.get('example_vi', '').strip()
        example_layout = example.get('example_layout', 'multiline').strip() or 'multiline'

        if example_layout == 'inline':
            inline_parts = []
            if example_cn:
                inline_parts.append(example_cn)
            if example_pinyin:
                inline_parts.append(f"/{example_pinyin}/")
            inline_text = "".join(inline_parts)
            if inline_text and example_vi:
                inline_text = f"{inline_text}: {example_vi}"
            elif example_vi:
                inline_text = example_vi

            if inline_text:
                p_example_inline = doc.add_paragraph()
                run_inline = p_example_inline.add_run(f"  - {inline_text}")
                if not example_cn and not example_pinyin:
                    run_inline.italic = True
            continue

        if example_cn:
            p_example_cn = doc.add_paragraph()
            p_example_cn.add_run(f"  - {example_cn}")

        if example_pinyin:
            p_example_py = doc.add_paragraph()
            p_example_py.add_run(f"    {example_pinyin}")

        if example_vi:
            p_example_vi = doc.add_paragraph()
            run_example_vi = p_example_vi.add_run(f"    {example_vi}")
            run_example_vi.italic = True


def _append_practice_block(doc, practice):
    practice = practice or {}
    title = practice.get('title', '').strip()
    content_lines = practice.get('content_lines', []) or []
    content_lines = [line.strip() for line in content_lines if str(line).strip()]

    if not title and not content_lines:
        return

    p_practice_title = doc.add_paragraph()
    run_practice_title = p_practice_title.add_run(title or "Luyện tập")
    run_practice_title.bold = True

    for line in content_lines:
        doc.add_paragraph(line)


def response2docx(file_path, prompt, file_name, project_id, creds, model_name):
    client = VertexClient(project_id, creds, model_name)

    with open(file_path, "rb") as f:
        pdf_data = f.read()
    mime_type = mimetypes.guess_type(file_path)[0] or "application/pdf"
    
    # Đọc Schema JSON
    schema_path = "resources/schema/summary_schema.json"
    if not os.path.exists(schema_path):
        raise FileNotFoundError(f"Không tìm thấy file schema tại {schema_path}")
        
    with open(schema_path, 'r', encoding='utf-8') as f:
        response_schema = json.load(f)

    # Gửi yêu cầu với JSON Schema
    AIresponse_text = client.send_data_to_AI(
        prompt,
        data=pdf_data, 
        mime_type=mime_type,
        response_mime_type="application/json",
        response_schema=response_schema,
        service_tier="auto"
    )
    print("Đã nhận phản hồi từ AI (JSON).")
    
    # Parse JSON
    try:
        data = json.loads(AIresponse_text)
    except json.JSONDecodeError as e:
        print(f"Lỗi phân tích JSON: {e}")
        print("Phản hồi thô:", AIresponse_text)
        raise ValueError("AI không trả về JSON hợp lệ.")
    
    doc = Document()
    
    # 1. Phần Mở đầu
    if "document_title_cn" in data and "document_title_vn" in data:
        p_title = doc.add_paragraph()
        run_title = p_title.add_run(f"Tóm tắt Bài học – {data['document_title_cn']} ({data['document_title_vn']})")
        run_title.bold = True
        
    if "lesson_summary" in data:
        doc.add_paragraph(data['lesson_summary'])
        doc.add_paragraph() # Dòng trống
        
    # 2. Phần Từ mới
    if data.get("vocabulary_list"):
        p_vocab_title = doc.add_paragraph()
        run_vocab_title = p_vocab_title.add_run("I. Từ Mới")
        run_vocab_title.bold = True
        
        for idx, vocab in enumerate(data['vocabulary_list'], 1):
            p = doc.add_paragraph()
            p.add_run(f"{idx}.  ") # Text thường
            
            run_hanzi = p.add_run(vocab.get('hanzi', ''))
            run_hanzi.bold = True # In đậm chữ Hán
            
            p.add_run(f" /{vocab.get('pinyin', '')}/ ({vocab.get('type', '')}): {vocab.get('meaning', '')}")
            
        doc.add_paragraph() # Dòng trống

    # 3. Phần Nội dung bài khóa
    if data.get("lessons"):
        p_lesson_title = doc.add_paragraph()
        run_lesson_title = p_lesson_title.add_run("II. Nội dung bài khóa")
        run_lesson_title.bold = True
        
        for idx, lesson in enumerate(data['lessons'], 1):
            p_lesson_name = doc.add_paragraph()
            run_lesson_name = p_lesson_name.add_run(f"Bài khóa {idx}")
            run_lesson_name.bold = True
            
            # Tiếng Trung
            if lesson.get('content_cn'):
                for line in lesson['content_cn'].split('\n'):
                    if line.strip():
                        doc.add_paragraph(line.strip())

                doc.add_paragraph()

            # Pinyin
            if lesson.get('content_pinyin'):
                for line in lesson['content_pinyin'].split('\n'):
                    if line.strip():
                        p_py = doc.add_paragraph()
                        run_py = p_py.add_run(line.strip())
                        run_py.italic = True

                doc.add_paragraph()
            
            # Tiếng Việt (In nghiêng)
            if lesson.get('content_vn'):
                for line in lesson['content_vn'].split('\n'):
                    if line.strip():
                        p_vn = doc.add_paragraph()
                        run_vn = p_vn.add_run(line.strip())
                        run_vn.italic = True
            
            # Lưu ý quan trọng
            notes = lesson.get('important_notes', [])
            if notes:
                p_note_title = doc.add_paragraph()
                p_note_title.add_run("* ").bold = False
                run_note_title = p_note_title.add_run("Lưu ý quan trọng:")
                run_note_title.bold = True
                
                for note in notes:
                    p_note = doc.add_paragraph()
                    p_note.add_run("-   ")

                    term_cn = note.get('term_cn', '').strip()
                    term_pinyin = note.get('term_pinyin', '').strip()
                    legacy_hanzi_pinyin = note.get('hanzi_pinyin', '').strip()
                    label_text = term_cn
                    if term_pinyin:
                        label_text = f"{term_cn} ({term_pinyin})"
                    elif not label_text:
                        label_text = legacy_hanzi_pinyin

                    run_note_cn = p_note.add_run(label_text)
                    run_note_cn.bold = True

                    p_note.add_run(f": {note.get('explanation', '')}")

                    _append_note_examples(doc, note.get('examples', []))

            _append_practice_block(doc, lesson.get('practice', {}))
                    
            doc.add_paragraph() # Dòng trống giữa các bài khóa

    output_dir = "output/summary"
    os.makedirs(output_dir, exist_ok=True)
    doc.save(os.path.join(output_dir, f"{file_name}.docx"))
    print(f"Đã lưu file Word tại: {output_dir}/{file_name}.docx")